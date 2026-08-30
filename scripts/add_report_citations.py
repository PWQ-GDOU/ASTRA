from docx import Document
p=r"H:\项目代码\航天器\ASTRA_技术报告_评委优化版.docx"
d=Document(p)
marks=[
 ('为此，本方案引入统一的健康指标',' [1, 6, 7]'),
 ('本场景选取NASA锂离子电池退化数据集',' [1, 2]'),
 ('本场景选取NASA锂离子电池退化数据集作为第一个迁移目标域',' [1, 2]'),
 ('由于真实在轨反作用轮遥测数据获取困难',' [3, 4, 5]'),
 ('目标场景一：NASA电池长期衰减',' [1, 2]'),
 ('目标场景二：反作用轮长期机械退化',' [3, 4, 5]'),
 ('趋势先验（target prior）的筛选隔离',' [6, 7, 9]'),
 ('辅助指标包括MAE、bias与归一化RMSE',' [9, 10]'),
 ('不确定性',' [8, 10, 11]'),
]
for para in d.paragraphs:
 t=para.text
 for prefix,cite in marks:
  if t.startswith(prefix) and cite.strip() not in t:
   para.add_run(cite)
   break
d.save(p)
print(p)
