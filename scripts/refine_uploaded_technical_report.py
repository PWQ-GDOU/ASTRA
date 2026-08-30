from pathlib import Path
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from docx.text.paragraph import Paragraph
ROOT=Path(r"H:\项目代码\航天器")
SRC=ROOT/"技术报告.docx"
OUT=ROOT/"ASTRA_技术报告_评委优化版.docx"
REPO="https://github.com/PWQ-GDOU/ASTRA"
def cc(p): return bool({n.tag.split('}')[-1] for n in p._p.iter()} & {'oMath','oMathPara','drawing','pict','object'})
def st(p,s):
 if not cc(p): p.text=s
def ins(a,s='',style=None):
 e=OxmlElement('w:p'); a._p.addnext(e); p=Paragraph(e,a._parent)
 if style: p.style=style
 if s:p.add_run(s)
 return p
def link(p,text,url):
 rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');rp.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def rm(p): p._element.getparent().remove(p._element)
def setup(d):
 n=d.styles['Normal'];n.font.name='宋体';n._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');n.font.size=Pt(10.5);n.paragraph_format.line_spacing=1.15;n.paragraph_format.space_after=Pt(5)
 for name,size,color in [('Heading 1',15,'1F4E79'),('Heading 2',12,'2F5597')]:
  s=d.styles[name];s.font.name='黑体';s._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string(color)
def classify(d):
 ch=re.compile(r'^[一二三四五六七八九十]+、');sub=re.compile(r'^\d+\.\d+')
 for p in d.paragraphs:
  t=' '.join(p.text.split())
  if not t:continue
  if t in {'摘要','关键词','结论','参考文献','评审材料获取与复现入口','附录A 交付物索引'} or ch.match(t):p.style='Heading 1'
  elif sub.match(t):p.style='Heading 2'
  elif t.startswith('表'):p.style='Normal'
def rew(d):
 rules=[
 ('航天器在轨服役期间，其关键部件的性能退化','航天器关键部件退化周期长、故障样本稀缺，真实在轨全寿命数据难以支撑直接建模。本项目以公开退化数据学习通用演化表征，再在航天电池与反作用轮仿真场景中进行小样本适配，并以可追溯协议记录数据、切分、训练和评估过程。'),
 ('我们以公开的喷嘴烧蚀退化数据作为源域知识库','源域采用喷嘴烧蚀多轨迹数据；目标域包括 NASA 电池和反作用轮退化仿真代理。目标侧采用轨迹级或轴承级 outer-fold 留出，比较 pretrained transfer、matched Scratch、Ridge 与趋势基线。COMSOL 结果属于反作用轮仿真代理的探索性证据，不外推为真实在轨遥测结论。'),
 ('值得强调的是，我们将NASA公开电池数据完整纳入迁移学习框架','NASA 电池链路完成了严格切分、目标侧尺度估计和逐折评估，但当前 N=1、N=2 均未达到预设的内部正向迁移标准，因此本报告将其作为诊断结果。所有主实验均输出逐折预测、指标、协议、数据清单和环境记录，支持复核。'),
 ('综上，本项目不仅验证了“用地面数据训练、为在轨场景服务”','本报告的主要方法学证据来自 FEMTO 轴承 -> COMSOL 反作用轮仿真代理链路的探索性复现；电池链路保留为未通过内部验收的诊断结果。最终是否满足赛事评分要求，由评委依据提交材料和现场复核判定。'),
 ('本报告严格对标赛事章程中的答题要求与评分体系','根据赛事方案，本文覆盖问题建模与仿真场景、数据集与仿真方法、PyTorch 可执行代码、迁移预测验证及工程表达。评审证据采用报告章节、仓库路径和机器可读结果文件三重索引。'),
 ('据此，本方案定位为“地面知识驱动、航天场景验证、工程闭环交付”','据此，本方案定位为“地面知识驱动、航天场景验证、工程闭环交付”。下表仅用于帮助评委定位官方评审要点与证据位置，不代表本项目自评得分或替代评委判断：'),
 ('综上，本方案在全部评分维度上均提供了可验证的证据支撑','后文分别给出对应的实验、代码、数据和审计文件；评审方可按仓库索引独立复核。'),
 ('项目自主构建了反作用轮长期机械退化的高保真退化模拟数据集','项目自主构建了反作用轮长期机械退化的参数化退化仿真数据集'),
 ('以上材料共同构成了从“算法验证”到“工程展示”的闭环交付','工程演示、启动命令和交付物索引见仓库 docs/engineering_demo.md；评审方可先运行工程演示，再按复现脚本检查协议和结果文件。'),
 ('为支持无需Python环境即可快速验证的可复现性目标','仓库提供 Dockerfile、docker/docker-compose.yml、docker/requirements.lock.txt 以及 reproduce_femto_ims_to_comsol.ps1/.sh。Docker 复现流程按 preflight -> smoke -> full -> artifact audit 执行，结果写入宿主机 outputs/，具体命令见 README.md 和 docs/docker_reproduction_report_20260821.md。'),
 ('以上结果共同证明：本项目代码库已在容器化层面实现了“一次构建，随处复现”','上述材料支持评审方在 Docker Desktop 中复现实验协议和工程演示；具体运行时间与机器环境可能不同，最终以生成的环境记录、协议文件和哈希清单为准。'),
 ('本项目不将仿真代理结论外推至真实在轨放行判定','本项目不将 COMSOL 仿真代理结论外推至真实在轨放行判定。所有“正向证据”均限定为协议内部的可审计结果；v5 明确标注为 exploratory_post_hoc_replication，仍需未见 outer holdout 或真实遥测进一步确认。')]
 for p in list(d.paragraphs):
  old=p.text
  for a,b in rules:
   if old.startswith(a):st(p,b);break
  if p.text.startswith('为全面刻画模型预测性能，项目采用多指标并行评估策略'):st(p,'项目同时报告 raw RMSE、MAE、bias 和按训练侧尺度归一化的 RMSE，以兼顾总体误差、典型误差和系统性偏差。')
  elif p.text.strip() in {'指标体系。','内部验收标准。'}:st(p,p.text.strip('。'))
  elif p.text.startswith('为将算法能力向工程可用性转化，项目在代码仓库之外配套开发了离线寿命状态仪表板'):rm(p)
def repo_block(d):
 a=next((p for p in d.paragraphs if p.text.strip()=='一、赛事要求与方案定位'),None)
 if not a:return
 h=ins(a,'评审材料获取与复现入口','Heading 1');p=ins(h);p.add_run('公开仓库：').bold=True;link(p,REPO,REPO);p.add_run('。GitHub 用于获取代码、数据清单、审计结果和复现说明；赛事正式作品仍以官方申报系统要求的作品包和报名表 PDF 为准。')
 p=ins(p,'推荐阅读顺序：README.md -> DATA_SOURCES.md -> docs/competition/competition_rules.pdf -> docs/engineering_demo.md -> outputs/ 中的协议和审计文件。')
 ins(p,'主实验入口为 scripts/exp_cross_transfer_v3.py；FEMTO -> COMSOL 外层留出入口为 scripts/exp_femto_ims_to_comsol_outerloo_v5.py；Docker 入口为 Dockerfile、docker/docker-compose.yml 和 docker/requirements.lock.txt。')
def software(d):
 a=next((p for p in d.paragraphs if p.text.strip()=='八、工程实现与可复现性'),None)
 if not a:return
 p=ins(a,'软件系统作为工程演示层','Heading 2')
 p=ins(p,'同学补充的软件部分被纳入工程演示层，而不是主实验层。软件系统读取已归档的预测结果和协议元数据，提供寿命趋势回放、RUL 数值、健康状态分级、时间轴浏览以及操作员/审计员双视图。它不重新训练模型，也不改变主实验的 outer-fold 切分、指标或结论，因此不会把展示功能误当成算法证据。')
 p=ins(p,'评审方可在仓库 outputs/engineering_demo/femto_ims_to_comsol_v5/ 查看可离线打开的 index.html 和 DASHBOARD_DATA.json；启动说明见 docs/engineering_demo.md，脚本入口为 scripts/run_engineering_demo.py 与 scripts/run_engineering_demo.ps1。该模块对应赛事“寿命状态直观表达和工程应用流程”要求，主实验仍由 scripts/exp_cross_transfer_v3.py 及其审计输出定义。')
 ins(p,'分支策略：main 只承载主实验、数据协议、复现环境和评委版文档；软件工程演示作为同一仓库中的独立模块维护，必要时使用 feature/software 或 software-demo 分支开发，合并前不得改写主实验结果。')
def refs(d):
 h=d.add_paragraph('参考文献',style='Heading 1');h.paragraph_format.page_break_before=True
 rs=['[1] Nguyen C. D., Bae S. J. Equivalent circuit simulated deep network architecture and transfer learning for remaining useful life prediction of lithium-ion batteries. Journal of Energy Storage, 2023, 71: 108042. 仓库：references/Equivalent circuit simulated deep network architecture and transfer.pdf.','[2] Hou G., Zhang F., Huang C., et al. Joint prediction of SOH and RUL for lithium-ion batteries by an enhanced Transformer model with physical information constraints. Energy, 2025, 336: 138435. 仓库：references/Joint prediction of SOH and RUL for Lithium-ion batteries by an enhanced.pdf.','[3] Muthuganapathy P., Chaturvedi S. K., Gargama H., et al. Friction torque prediction of precision ball bearing unit for reaction wheel actuators for spacecraft applications. International Journal of System Assurance Engineering and Management, 2025. 仓库：references/Friction torque prediction of precision ball bearing unit for reaction wheel actuators for spacecraft applications.pdf.','[4] Alidadi M., Rahimi A. Fault Diagnosis of Lubrication Decay in Reaction Wheels Using Temperature Estimation and Forecasting via Enhanced Adaptive Particle Filter. Sensors. 仓库：references/Fault Diagnosis of Lubrication Decay in Reaction Wheels Using Temperature Estimation and Forecasting via Enhanced Adaptive Particle Filter.pdf.','[5] Yücesan Y. A., Viana F. A. C. Physics-informed digital twin for wind turbine main bearing fatigue: Quantifying uncertainty in grease degradation. Applied Soft Computing, 2023, 149: 110921. 仓库：references/Physics informed digital twin for wind turbine main bearing fatigue Quantifying uncertainty in grease degradation.pdf.','[6] Shang J., Xu D., Qiu H., et al. Domain generalization for rotating machinery real-time remaining useful life prediction via multi-domain orthogonal degradation feature exploration. Mechanical Systems and Signal Processing, 2025, 223: 111924. 仓库：references/Domain Generalization for Rotating Machinery Real-Time Remaining Useful Life Prediction via Multi-Domain Orthogonal Degradation Feature Exploration.pdf.','[7] Nejjar I., Geissmann F., Zhao M., et al. Domain adaptation via alignment of operation profile for remaining useful lifetime prediction. Reliability Engineering and System Safety. 仓库：references/Domain Adaptation via Alignment of Operation Profile for Remaining Useful Lifetime Prediction.pdf.','[8] Chang L., Lin Y.-H. Few-shot remaining useful life prediction based on Bayesian meta-learning with predictive uncertainty calibration. Engineering Applications of Artificial Intelligence, 2025, 142: 109980. 仓库：references/Few-Shot Remaining Useful Life Prediction Based on Bayesian Meta-Learning with Predictive Uncertainty Calibration.pdf.','[9] Lin W., Chen X., Lu H., et al. A novel interactive prognosis framework with nonlinear Wiener process and multi-sensor fusion for remaining useful life prediction. Journal of Process Control, 2024, 140: 103264. 仓库：references/A Novel Interactive Prognosis Framework with Nonlinear Wiener Process and Multi-Sensor Fusion for Remaining Useful Life Prediction.pdf.','[10] Xiang F., Zhang Y., Zhang S., et al. Bayesian gated-transformer model for risk-aware prediction of aero-engine remaining useful life. Expert Systems with Applications, 2024, 238: 121859. 仓库：references/Bayesian Gated-Transformer Model for Risk-Aware Prediction of Aero-Engine Remaining Useful Life.pdf.','[11] Wang E., Lei Z., Wen G., et al. A physics-constrained Bayesian neural network for machinery remaining useful life prediction and uncertainty quantification. Reliability Engineering and System Safety, 2026, 266: 111778. 仓库：references/A Physics-Constrained Bayesian Neural Network for Machinery Remaining Useful Life Prediction and Uncertainty Quantification.pdf.','[12] 本项目方法相关参考：Self-Supervised Health Representation Decomposition Based on Contrast Learning；Pre-training Enhanced Unsupervised Contrastive Domain Adaptation for Industrial Equipment Remaining Useful Life Prediction；Joint Domain-Adaptive Transformer Model for Bearing RUL Prediction Across Different Domains。均存于仓库 references/，用于方法设计背景与对照，不替代本文独立实验。']
 for x in rs:
  p=d.add_paragraph(x);p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5);p.paragraph_format.space_after=Pt(3)
  for r in p.runs:r.font.size=Pt(9.5)
def index(d):
 h=d.add_paragraph('附录A 交付物索引',style='Heading 1');h.paragraph_format.page_break_before=True;p=d.add_paragraph();p.add_run('仓库地址：').bold=True;link(p,REPO,REPO)
 xs=[('赛事规则原件','docs/competition/competition_rules.pdf'),('数据与来源说明','DATA_SOURCES.md；data/raw/competition/；third_party/'),('主跨组件迁移协议','scripts/exp_cross_transfer_v3.py；outputs/cross_component_transfer_v3/'),('FEMTO -> COMSOL outer-fold 协议','scripts/exp_femto_ims_to_comsol_outerloo_v5.py；outputs/femto_ims_to_comsol_outerloo_v5_exploratory_fitonly_innerloo_full5x120_20260825/'),('软件工程演示','scripts/run_engineering_demo.py；scripts/run_engineering_demo.ps1；docs/engineering_demo.md；outputs/engineering_demo/femto_ims_to_comsol_v5/'),('Docker 与锁定环境','Dockerfile；docker/docker-compose.yml；docker/requirements.lock.txt'),('一键复现','scripts/reproduce_femto_ims_to_comsol.ps1；scripts/reproduce_femto_ims_to_comsol.sh；README.md'),('审计文件','PROTOCOL.json；DATA_MANIFEST.json；PREFLIGHT.json；ACCEPTANCE.json；OUTER_FOLD_SUMMARY.csv；PREDICTIONS.csv')]
 for a,b in xs:
  p=d.add_paragraph(style='Normal');p.add_run(a+'：').bold=True;p.add_run(b)
 d.add_paragraph('说明：GitHub 仓库用于技术材料获取与复核；赛事正式提交仍应遵循官方申报系统、作品压缩包和报名表 PDF 的要求。')
def main():
 d=Document(str(SRC));setup(d);rew(d);update_tables(d);repo_block(d);software(d);classify(d);refs(d);index(d);d.core_properties.title='ASTRA 航天器关键组件跨域退化迁移寿命预测技术报告（评委优化版）';d.core_properties.subject='赛事技术方案、数据仿真、迁移验证、软件演示与复现材料索引';d.core_properties.author='ASTRA 项目组';d.save(str(OUT));print(OUT)
def update_tables(d):
 if not d.tables:return
 t=d.tables[0]
 if t.cell(0,0).text.strip()=='评分维度':t.cell(0,0).text='官方评审维度';t.cell(0,1).text='官方权重';t.cell(0,2).text='官方要求与本项目证据位置（不代表自评得分）'
 for table in d.tables:
  for row in table.rows:
   for cell in row.cells:
    old=cell.text; new=old.replace('高度结构相似','分别说明退化机理与可观测量').replace('高保真退化模拟数据','参数化退化仿真数据')
    if new != old: cell.text=new
if __name__=='__main__':main()
